#!/usr/bin/env python3
"""Dal foglio delle acque integrate al file che il server sovrappone al catalogo.

    python3 genera_integrazioni.py Acque_dati_integrati_WikiAcqua.xlsx [catalogo.json] [uscita.json]

- integra: i valori che il database ha vuoti e il foglio ha compilato
- correggi: gli errori certi del foglio "Verifica sospetti" (vedi CORREZIONI);
            ognuna vale solo finche' il database ha ancora il valore sbagliato
- nuove:   le acque del foglio assenti dal database, con numeri di scheda che
           non possono scontrarsi con quelli futuri del cliente

Sparea resta com'e': serve l'etichetta. Montoso e Fonte Azzurrina sono
svuotate (valore null) come consiglia il foglio.
"""
import datetime
import json
import os
import re
import sys

import openpyxl

QUI = os.path.dirname(os.path.abspath(__file__))
XLSX = sys.argv[1]
CATALOGO = sys.argv[2] if len(sys.argv) > 2 else os.path.join(QUI, "fixtures", "catalogo_attuale.json")
USCITA = sys.argv[3] if len(sys.argv) > 3 else os.path.join(QUI, "data", "integrazioni.json")

COLONNE = {
    "Residuo fisso a 180 °C": "RESIDUO_FISSO_A_180", "Sodio": "SODIO", "Calcio": "CALCIO",
    "Magnesio": "MAGNESIO", "Bicarbonati": "BICARBONATI", "Nitrati": "NITRATI",
    "Potassio": "POTASSIO", "Solfati": "SOLFATI", "Cloruri": "CLORURI", "Fluoruri": "FLUORURI",
    "Silice": "SILICE", "pH alla sorgente": "PH_ALLA_SORGENTE",
    "Temperatura alla sorgente (°C)": "TEMP_ALLA_SORGENTE",
    "Anidride carbonica libera": "ANIDRIDE_CARBONICA_LIBERA", "Stronzio": "STRONZIO",
    "Ossidabilità": "OSSIDABILITA",
}

# Gli errori certi del foglio "Verifica sospetti" con un valore proposto
# preciso: (id, campo, valore sbagliato nel database, valore giusto, perche').
CORREZIONI = [
    (2, "PH_ALLA_SORGENTE", 1.27, 5.86, "etichetta 13/11/2023 (Beverfood)"),
    (58, "TEMP_ALLA_SORGENTE", 170, 17.0, "virgola persa"),
    (66, "SODIO", 0.01, 14.8, "valore scritto in g/L; etichetta del produttore"),
    (67, "RESIDUO_FISSO_A_180", 8.37, 275.4, "analisi trovata"),
    (97, "RESIDUO_FISSO_A_180", 0.56, 560, "valore scritto in g/L"),
    (114, "PH_ALLA_SORGENTE", 2.5, 7.7, "era l'anidride carbonica finita nella colonna del pH"),
    (208, "RESIDUO_FISSO_A_180", 7.4, 690, "il pH era finito nella colonna del residuo fisso"),
    (208, "PH_ALLA_SORGENTE", 2.9, 7.4, "il pH era finito nella colonna del residuo fisso"),
    (216, "SODIO", 0.01, 13, "valore scritto in g/L"),
    # Le cinque schede rimaste incoerenti. Nessuna ha un'analisi in testo online:
    # si corregge solo cio' che i numeri stessi dimostrano.
    # SANTA CHIANCIANO: residuo fisso e solfati scritti in g/L e calcio con la
    # virgola spostata. Con Ca 738 e SO4 1750 lo squilibrio ionico passa dal 21%
    # al 4%, il residuo fisso atteso (3162) coincide con 3200 e con i 3280 di
    # Beverfood. I valori sono dell'analisi in uso nel database, non della storica.
    (14, "RESIDUO_FISSO_A_180", 3.2, 3200, "scritto in g/L (Beverfood: 3280)"),
    (14, "CALCIO", 73.8, 738, "virgola spostata: bilancio ionico 21% -> 4%"),
    (14, "SOLFATI", 1.75, 1750, "scritto in g/L (Terme di Chianciano: fino a circa 1800)"),
    # NOVELLA: 0,40 e' g/L. Con le altre analisi il residuo fisso atteso e' >= 342.
    (159, "RESIDUO_FISSO_A_180", 0.4, 400, "scritto in g/L; da confermare con l'etichetta"),
    # SOLARIA: 0,01 mg/L di sodio non esiste (in altre schede era 0,01 g/L, cioe'
    # circa 10-15 mg/L), ma il valore vero non si trova: si toglie.
    (218, "SODIO", 0.01, None, "valore in g/L, quello vero non si trova online"),
    # CONTESSA e TOSCA: gli stessi otto numeri tondi, identici fra loro, sono
    # segnaposto. Si tolgono; restano residuo fisso e pH.
    *[(i, c, v, None, "segnaposto identico in Contessa e Tosca")
      for i in (55, 234)
      for c, v in (("SODIO", 200), ("CALCIO", 150), ("MAGNESIO", 50), ("BICARBONATI", 600),
                   ("SOLFATI", 200), ("CLORURI", 200), ("NITRATI", 10), ("FLUORURI", 4))],
    # Schede con numeri segnaposto: il foglio consiglia di svuotarle finche'
    # non si ha l'etichetta. Vuote non entrano nelle classifiche per sbaglio.
    *[(74, c, 0.1, None, "FONTE AZZURRINA: sette valori tutti 0,1, segnaposto")
      for c in ("SODIO", "CALCIO", "MAGNESIO", "POTASSIO", "SOLFATI", "CLORURI", "NITRATI")],
    (147, "RESIDUO_FISSO_A_180", 244, None, "MONTOSO: scheda corrotta, da rifare dall'etichetta"),
    (147, "SODIO", 150, None, "MONTOSO: scheda corrotta, da rifare dall'etichetta"),
    (147, "SOLFATI", 490, None, "MONTOSO: scheda corrotta, da rifare dall'etichetta"),
    (147, "CLORURI", 26, None, "MONTOSO: scheda corrotta, da rifare dall'etichetta"),
    (147, "TEMP_ALLA_SORGENTE", 76, None, "MONTOSO: scheda corrotta, da rifare dall'etichetta"),
    (147, "ANIDRIDE_CARBONICA_LIBERA", 100, None, "MONTOSO: scheda corrotta, da rifare dall'etichetta"),
]

PRIMA_ID_NUOVA = 90001      # il database del cliente non arrivera' mai qui


def numero(v):
    if v is None or str(v).strip().upper() in ("", "MANCA", "NULL"):
        return None
    return float(str(v).replace(",", "."))


def testo(v):
    return f"{v:.2f}"


def classificazione(rf):
    """Stesse soglie di legge usate dalla scheda: minimamente mineralizzata
    fino a 50 mg/L, oligominerale fino a 500, minerale fino a 1500, poi ricca di sali."""
    if rf is None:
        return None
    if rf <= 50:
        return "Min. Mineralizzata"
    if rf <= 500:
        return "Oligominerale"
    if rf <= 1500:
        return "Minerale"
    return "Ricca di Sali"


def vuoto(v):
    return v is None or str(v).strip().lower() in ("", "null")


db = {int(x["ID"]): x for x in json.load(open(CATALOGO, encoding="utf-8"))}
modello = {k: None for k in next(iter(db.values()))}
wb = openpyxl.load_workbook(XLSX, data_only=True)
righe = list(wb["Tutte le acque"].iter_rows(values_only=True))
foglio = {}
for r in righe[1:]:
    if r[0] is not None:
        foglio[int(r[0])] = dict(zip(righe[0], r))

integra = {}
for i, d in sorted(foglio.items()):
    if i not in db:
        continue
    nuovi = {}
    for col, campo in COLONNE.items():
        v = numero(d[col])
        if v is not None and vuoto(db[i].get(campo)):
            nuovi[campo] = testo(v)
    if nuovi:
        integra[str(i)] = nuovi

correggi = []
for i, campo, sbagliato, giusto, motivo in CORREZIONI:
    attuale = float(db[i][campo])
    assert abs(attuale - sbagliato) < 1e-9, f"{i} {campo}: nel database c'e' {attuale}, non {sbagliato}"
    correggi.append({"id": str(i), "nome": db[i]["NOME_ETICHETTA"].strip(), "campo": campo,
                     "se_ancora": sbagliato, "valore": giusto, "motivo": motivo})

note = {}
for r in list(wb["Acque aggiunte"].iter_rows(values_only=True))[1:]:
    if r[0] is not None:
        note[int(r[0])] = {"fonte": r[4], "link": r[5], "note": r[6]}

nuove, fonti = [], {}
for i, d in sorted(foglio.items()):
    if i in db:
        continue
    scheda = dict(modello)
    nuovo_id = PRIMA_ID_NUOVA + (i - min(k for k in foglio if k not in db))
    scheda.update(ID=nuovo_id, NOME_ETICHETTA=d["Acqua"], REGIONE=d["Regione"])
    for col, campo in COLONNE.items():
        v = numero(d[col])
        scheda[campo] = None if v is None else testo(v)
    scheda["CLASSIFICAZIONE"] = classificazione(numero(d["Residuo fisso a 180 °C"]))
    m = re.search(r"[Ff]onte a ([^.;]+?)\s*(?:[.;]|$)", str(note.get(i, {}).get("note") or ""))
    scheda["STABILIMENTO_1"] = m.group(1).strip() if m else None
    nuove.append(scheda)
    fonti[str(nuovo_id)] = {"scheda_foglio": i, **note.get(i, {})}

os.makedirs(os.path.dirname(USCITA), exist_ok=True)
with open(USCITA, "w", encoding="utf-8") as f:
    json.dump({"generato": datetime.date.today().isoformat(), "origine": os.path.basename(XLSX),
               "integra": integra, "correggi": correggi, "nuove": nuove, "fonti_nuove": fonti},
              f, ensure_ascii=False, indent=1)
print(f"integra: {sum(len(v) for v in integra.values())} valori in {len(integra)} acque | "
      f"correggi: {len(correggi)} | nuove: {len(nuove)} -> {USCITA}")
