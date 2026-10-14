#!/usr/bin/env python3
"""Genera il foglio con le acque a cui mancano dati, da inoltrare al cliente.

    python3 rapporto_dati_mancanti.py [catalogo.json] [uscita.xlsx]
"""
import asyncio, datetime, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

QUI = os.path.dirname(os.path.abspath(__file__))
FILE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(QUI, "fixtures", "catalogo_attuale.json")
USCITA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(QUI, "Acque_dati_mancanti_WikiAcqua.xlsx")
os.environ.update(GEMINI_API_KEY="x", DB_ACQUE_FILE=FILE)
sys.path.insert(0, os.path.join(QUI, "app"))
import chat                                                    # per riusare i controlli di coerenza

grezzo = json.load(open(FILE, encoding="utf-8"))
schede = {s["id"]: s for s in asyncio.run(chat.catalogo())}   # senza doppioni, con le incoerenze
doppioni = [x for x in grezzo if str(x["ID"]) not in schede]
righe = [x for x in grezzo if str(x["ID"]) in schede]

# (chiave nel database, nome, importanza)
ALTA = "Alta: decide le categorie d'uso dell'app"
MEDIA = "Media"
BASSA = "Bassa: quasi mai in etichetta"
PROP = [("RESIDUO_FISSO_A_180", "Residuo fisso a 180 °C", ALTA), ("SODIO", "Sodio", ALTA), ("CALCIO", "Calcio", ALTA),
        ("MAGNESIO", "Magnesio", ALTA), ("BICARBONATI", "Bicarbonati", ALTA), ("NITRATI", "Nitrati", ALTA),
        ("POTASSIO", "Potassio", MEDIA), ("SOLFATI", "Solfati", MEDIA), ("CLORURI", "Cloruri", MEDIA),
        ("FLUORURI", "Fluoruri", MEDIA), ("SILICE", "Silice", MEDIA), ("PH_ALLA_SORGENTE", "pH alla sorgente", MEDIA),
        ("TEMP_ALLA_SORGENTE", "Temperatura alla sorgente (°C)", MEDIA),
        ("ANIDRIDE_CARBONICA_LIBERA", "Anidride carbonica libera", MEDIA),
        ("STRONZIO", "Stronzio", BASSA), ("OSSIDABILITA", "Ossidabilità", BASSA)]
ALTE = [p for p in PROP if p[2] == ALTA]


def manca(x, k):
    v = x.get(k)
    return v is None or str(v).strip() == "" or str(v).strip().lower() == "null"


def pulito(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return v


ARIAL = "Arial"
TESTA = PatternFill("solid", fgColor="D9D9D9")
ROSSO = PatternFill("solid", fgColor="F8CBAD")
GIALLO = PatternFill("solid", fgColor="FFF2CC")
sottile = Side(style="thin", color="BBBBBB")
BOX = Border(left=sottile, right=sottile, top=sottile, bottom=sottile)

wb = Workbook()


def foglio(titolo, intestazioni, larghezze, primo=False):
    ws = wb.active if primo else wb.create_sheet()
    ws.title = titolo
    for i, (h, w) in enumerate(zip(intestazioni, larghezze), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=ARIAL, bold=True, size=10); c.fill = TESTA; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def scrivi(ws, r, valori, wrap_da=None):
    for i, v in enumerate(valori, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=ARIAL, size=10); c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------- Riepilogo
n = len(righe)
senza_alte = [x for x in righe if any(manca(x, k) for k, _, _ in ALTE)]
senza_quattro = [x for x in righe if any(manca(x, k) for k, _, _ in PROP[:4])]
complete = [x for x in righe if not any(manca(x, k) for k, _, _ in PROP)]

ws = wb.active
ws.title = "Riepilogo"
ws["A1"] = "WikiAcqua · dati mancanti nel catalogo"
ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
ws["A2"] = f"Catalogo letto da wikiacqua.it il {datetime.date.today().strftime('%d/%m/%Y')} · {n} acque"
ws["A2"].font = Font(name=ARIAL, size=10, color="666666")
testo = [
    ("Acque con tutte le 16 proprietà", len(complete)),
    ("Acque a cui manca almeno una delle 6 proprietà ad alta importanza", len(senza_alte)),
    ("Acque a cui manca almeno una delle 4 citate dal cliente (residuo fisso, sodio, calcio, magnesio)", len(senza_quattro)),
]
for i, (t, v) in enumerate(testo, 4):
    ws.cell(row=i, column=1, value=t).font = Font(name=ARIAL, size=10)
    ws.cell(row=i, column=2, value=v).font = Font(name=ARIAL, bold=True, size=10)
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws["A8"] = ("Perché conta: l'assistente classifica le acque (neonati, sportivi, ossa, poco sodio, digestive, leggere) con le 6 proprietà "
            "ad alta importanza. Un'acqua a cui ne manca una non può rientrare in quella categoria. Per i neonati i nitrati "
            "mancanti sono il dato più delicato.")
ws["A8"].font = Font(name=ARIAL, size=10); ws["A8"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A8:D8"); ws.row_dimensions[8].height = 52
for i, h in enumerate(["Proprietà", "Importanza", "Acque senza il dato", "% sul totale"], 1):
    c = ws.cell(row=10, column=i, value=h)
    c.font = Font(name=ARIAL, bold=True, size=10); c.fill = TESTA; c.border = BOX
for r, (k, nome, imp) in enumerate(PROP, 11):
    mancanti = sum(manca(x, k) for x in righe)
    scrivi(ws, r, [nome, imp, mancanti, mancanti / n])
    ws.cell(row=r, column=4).number_format = "0%"
    if imp == ALTA and mancanti:
        for col in range(1, 5):
            ws.cell(row=r, column=col).fill = GIALLO
for col, w in zip("ABCD", (62, 40, 20, 14)):
    ws.column_dimensions[col].width = w
ws.row_dimensions[4].height = ws.row_dimensions[5].height = 18
ws.row_dimensions[6].height = 30

# ----------------------------------------------------------- Da completare
ws = foglio("Da completare", ["ID", "Acqua", "Regione", "Proprietà importanti mancanti", "N.", "Altre proprietà mancanti", "Sito web"],
            [7, 34, 18, 52, 6, 56, 30])
ordine = sorted(senza_alte, key=lambda x: (-sum(manca(x, k) for k, _, _ in ALTE), x["NOME_ETICHETTA"].strip()))
for r, x in enumerate(ordine, 2):
    importanti = [nome for k, nome, _ in ALTE if manca(x, k)]
    altre = [nome for k, nome, imp in PROP if imp != ALTA and manca(x, k)]
    scrivi(ws, r, [x["ID"], x["NOME_ETICHETTA"].strip(), x["REGIONE"], ", ".join(importanti), len(importanti),
                   ", ".join(altre), x["SITO_WEB"] or ""])
ws.auto_filter.ref = f"A1:G{len(ordine) + 1}"

# ------------------------------------------------------------ Tutte le acque
ws = foglio("Tutte le acque", ["ID", "Acqua", "Regione"] + [p[1] for p in PROP] + ["N. mancanti"],
            [7, 34, 16] + [13] * len(PROP) + [11])
for r, x in enumerate(sorted(righe, key=lambda x: x["NOME_ETICHETTA"].strip()), 2):
    celle = [x["ID"], x["NOME_ETICHETTA"].strip(), x["REGIONE"]]
    assenti = 0
    for k, _, _ in PROP:
        if manca(x, k):
            celle.append("MANCA"); assenti += 1
        else:
            celle.append(pulito(x[k]))
    scrivi(ws, r, celle + [assenti])
    for c in range(4, 4 + len(PROP)):
        if ws.cell(row=r, column=c).value == "MANCA":
            ws.cell(row=r, column=c).fill = ROSSO
ws.auto_filter.ref = f"A1:{get_column_letter(3 + len(PROP) + 1)}{len(righe) + 1}"
ws.freeze_panes = "D2"

# --------------------------------------------------------- Valori sospetti
ws = foglio("Valori sospetti", ["ID", "Acqua", "Cosa non torna", "Residuo fisso", "Calcio", "Magnesio", "Sodio", "Potassio", "pH", "Temperatura"],
            [7, 34, 70, 11, 9, 10, 9, 10, 7, 12])
sospette = sorted((s for s in schede.values() if s["incoerenze"]), key=lambda s: s["nome"])
r = 2
for s in sospette:
    scrivi(ws, r, [s["id"], s["nome"], "; ".join(s["incoerenze"]), s["residuo_fisso_a_180"], s["calcio"], s["magnesio"],
                   s["sodio"], s["potassio"], s["ph_alla_sorgente"], s["temp_alla_sorgente"]])
    r += 1
for x in doppioni:
    originale = next(s for s in schede.values() if chat._chiave(s["nome"]) == chat._chiave(x["NOME_ETICHETTA"]))
    scrivi(ws, r, [x["ID"], x["NOME_ETICHETTA"].strip(),
                   f"Doppione dell'acqua ID {originale['id']} ({originale['nome']}): stesso stabilimento, stessi valori. Va tenuta una sola scheda.",
                   None, None, None, None, None, None, None])
    for c in range(1, 11):
        ws.cell(row=r, column=c).fill = GIALLO
    r += 1
ws.cell(row=r + 1, column=1, value=("Controllo automatico: residuo fisso sotto 5 mg/L; calcio+magnesio+sodio+potassio che superano il residuo fisso "
                                     "(fisicamente impossibile); pH fuori da 4-10; temperatura fuori da 0-100 °C; sodio sotto 0,05 mg/L. "
                                     "Sono segnalazioni da confrontare con l'etichetta, non verdetti.")).font = Font(name=ARIAL, size=9, italic=True)

wb.save(USCITA)
print(f"scritto: {USCITA}")
print(f"acque: {n} | da completare: {len(ordine)} | con tutte le proprietà: {len(complete)} | sospette: {len(sospette)} | doppioni: {len(doppioni)}")
